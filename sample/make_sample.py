"""產生「範例」Excel：結構與正式輸出相同，但數字是隨機假資料，只用來確認版面。"""
import os
import random
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "PortfolioData_W_20260803_20260831_SAMPLE.xlsx")

# (代號, 名稱, 類別, 幣別, 單位, 起始價)
ASSETS = [
    ("^TWII", "TSEC weighted index", "台灣加權指數 Taiwan Index", "TWD", "price", 43000),
    ("^N225", "Nikkei 225", "日經 225 Nikkei 225", "JPY", "price", 40000),
    ("SPY", "SPDR S&P 500 ETF Trust", "S&P 500", "USD", "price", 745),
    ("VEU", "Vanguard FTSE All-World ex-US ETF", "美國以外成熟市場 Developed ex-US", "USD", "price", 70),
    ("VWO", "Vanguard FTSE Emerging Markets ETF", "新興市場 Emerging Markets", "USD", "price", 55),
    ("LQD", "iShares iBoxx $ Investment Grade Corporate Bond ETF", "美國投資等級公司債 US IG Corporate Bond", "USD", "price", 110),
    ("VNQ", "Vanguard Real Estate ETF", "美國房地產 US Real Estate", "USD", "price", 95),
    ("GLD", "SPDR Gold Shares", "黃金 Gold", "USD", "price", 330),
    ("DBC", "Invesco DB Commodity Index Tracking Fund", "大宗商品 Commodities", "USD", "price", 24),
    ("^TNX", "CBOE Interest Rate 10 Year T Note", "美國 10 年期公債殖利率 US 10Y Yield", "USD", "yield_pct", 4.7),
    ("2330.TW", "Taiwan Semiconductor Manufacturing", "個股 Stock", "TWD", "price", 2400),
    ("MMM", "3M Company", "個股 Stock", "USD", "price", 150),
    ("AAPL", "Apple Inc.", "個股 Stock", "USD", "price", 230),
    ("JPM", "JPMorgan Chase & Co.", "個股 Stock", "USD", "price", 290),
    ("LVMUY", "LVMH Moet Hennessy Louis Vuitton (ADR)", "個股 Stock", "USD", "price", 120),
]

# (對齊日 = 週五, 預設實際最後交易日, 註記)
WEEKS = [
    (date(2026, 7, 31), date(2026, 7, 31), "Prior period 前一期"),
    (date(2026, 8, 7), date(2026, 8, 7), ""),
    (date(2026, 8, 14), date(2026, 8, 14), ""),
    (date(2026, 8, 21), date(2026, 8, 21), ""),
    (date(2026, 8, 28), date(2026, 8, 28), ""),
    (date(2026, 9, 4), date(2026, 8, 31), "Partial 未完整"),
]
# 示範：某些市場該週的實際最後交易日不同，或整週休市（Combined 會留白）
OVERRIDE_DATE = {("^TWII", date(2026, 8, 21)): date(2026, 8, 20)}
SKIP = {("^N225", date(2026, 8, 14))}

HEAD_FILL = PatternFill("solid", fgColor="DCE6F1")
SAMPLE_FILL = PatternFill("solid", fgColor="FFF2CC")
BOLD = Font(bold=True)


def build_rows(seed_i, sym, base, unit):
    rnd = random.Random(seed_i)
    rows, price = [], float(base)
    for key, d, flag in WEEKS:
        # 先產生亂數再決定是否略過，這樣略過某週不會改變其他週的數字
        step = rnd.uniform(-0.03, 0.035)
        vol = 0 if sym.startswith("^") else rnd.randint(2_000_000, 90_000_000)
        if unit == "price":
            price *= 1 + step
        else:
            price = price + step * 2
        if (sym, key) in SKIP:
            continue
        close = round(price, 2 if unit == "price" else 3)
        adj = close if unit == "yield_pct" else round(close * 0.985, 2)  # 假裝有配息
        rows.append((OVERRIDE_DATE.get((sym, key), d), key, close, adj, vol, flag))
    return rows


def style_header(ws, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.font, cell.fill = BOLD, HEAD_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "B2"
    ws.row_dimensions[1].height = 32


def main():
    wb = Workbook()
    info = wb.active
    info.title = "Info 說明"
    cadj = wb.create_sheet("Combined_AdjClose")
    cclose = wb.create_sheet("Combined_Close")

    data = {}
    for i, (sym, name, cat, cur, unit, base) in enumerate(ASSETS):
        data[sym] = build_rows(i, sym, base, unit)
        ws = wb.create_sheet(sym)
        if unit == "yield_pct":
            ws.append(["Date 日期", "Yield (%) 殖利率", "Adj Yield (%) 調整後殖利率", "Volume 成交量", "Flag 註記"])
        else:
            ws.append(["Date 日期", "Close 收盤價", "Adj Close 調整後收盤價", "Volume 成交量", "Flag 註記"])
        for d, _k, close, adj, vol, flag in data[sym]:
            ws.append([datetime(d.year, d.month, d.day), close, adj, vol, flag])
        for r in range(2, ws.max_row + 1):
            ws.cell(r, 1).number_format = "yyyy-mm-dd"
            ws.cell(r, 2).number_format = ws.cell(r, 3).number_format = "#,##0.00##"
            ws.cell(r, 4).number_format = "#,##0"
        style_header(ws, 5)
        for col, w in zip("ABCDE", (13, 16, 22, 16, 20)):
            ws.column_dimensions[col].width = w

    # Combined：外部聯集，缺值留白
    keys = [k for k, _d, _f in WEEKS]
    syms = [a[0] for a in ASSETS]
    for ws, idx in ((cadj, 3), (cclose, 2)):
        ws.append(["Date 日期"] + syms + ["Flag 註記"])
        for k in keys:
            line, flags = [datetime(k.year, k.month, k.day)], []
            for s in syms:
                hit = next((r for r in data[s] if r[1] == k), None)
                line.append(hit[idx] if hit else None)
                if hit and hit[5]:
                    flags.append(hit[5])
            line.append(flags[0] if flags else "")
            ws.append(line)
        for r in range(2, ws.max_row + 1):
            ws.cell(r, 1).number_format = "yyyy-mm-dd"
            for c in range(2, len(syms) + 2):
                ws.cell(r, c).number_format = "#,##0.00##"
        style_header(ws, len(syms) + 2)
        ws.column_dimensions["A"].width = 13
        for c in range(2, len(syms) + 2):
            ws.column_dimensions[get_column_letter(c)].width = 11
        ws.column_dimensions[get_column_letter(len(syms) + 2)].width = 20

    # Info
    info.append(["代號 Ticker", "名稱 Name", "類別 Category", "幣別 Currency", "單位 Unit",
                 "第一筆日期 First Date", "最後一筆日期 Last Date", "筆數 Rows", "備註 Notes"])
    for sym, name, cat, cur, unit, _b in ASSETS:
        rows = data[sym]
        note = "殖利率 %，非價格，不可直接計算報酬 / Yield in %, not a price" if unit == "yield_pct" else ""
        if sym == "^N225":
            note = "示範：2026-08-14 該週整週休市，Combined 留白 / demo: whole week closed, blank in Combined"
        if sym == "^TWII":
            note = "示範：2026-08-21 該週最後交易日為 08-20 / demo: last trading day of the week is 08-20"
        info.append([sym, name, cat, cur, "殖利率 % Yield" if unit == "yield_pct" else "價格 Price",
                     datetime.combine(rows[0][0], datetime.min.time()),
                     datetime.combine(rows[-1][0], datetime.min.time()), len(rows), note])
    style_header(info, 9)
    info.freeze_panes = "A2"
    for r in range(2, info.max_row + 1):
        info.cell(r, 6).number_format = info.cell(r, 7).number_format = "yyyy-mm-dd"
    for col, w in zip("ABCDEFGHI", (12, 46, 36, 12, 16, 18, 18, 10, 70)):
        info.column_dimensions[col].width = w

    notes = [
        ("【這是範例檔 SAMPLE】數字是隨機產生的假資料，只用來確認版面，不可用於報告。",
         "THIS IS A SAMPLE FILE. Numbers are random dummy data, for layout review only."),
        ("資料來源：Yahoo Finance；下載時間：2026-10-08 12:00（台北時間）",
         "Source: Yahoo Finance; downloaded at 2026-10-08 12:00 (Taipei time)"),
        ("頻率：週 Weekly；期間：2026-08-03 ~ 2026-08-31（往回推：一個月）",
         "Frequency: Weekly; period: 2026-08-03 ~ 2026-08-31 (lookback: one month)"),
        ("Close = 未調整股利（已調整分割）；Adj Close = 調整分割與股利",
         "Close = not adjusted for dividends (split-adjusted); Adj Close = adjusted for splits and dividends"),
        ("幣別：各標的維持當地幣別，未換匯",
         "Currency: each series stays in its local currency; no FX conversion"),
        ("^TNX 為殖利率（%），不是價格，不可直接計算報酬",
         "^TNX is a yield in %, not a price; do not compute returns from it directly"),
        ("日期標示：各標的分頁的週／月／年資料，一律標該期「實際最後交易日」",
         "Dates: weekly/monthly/yearly rows in each ticker sheet use the actual last trading day of the period"),
        ("Combined 對齊：日資料以交易日；週資料以該週週五；月資料以月底；年資料以 12/31",
         "Combined alignment: daily = trading day; weekly = that week's Friday; monthly = month end; yearly = Dec 31"),
        ("Combined 留白：該期該市場沒有交易時，儲存格為空白，不補值。計算共變異數前請自行處理",
         "Combined blanks: a cell is empty when that market has no data for the period; handle before computing covariances"),
        ("Flag 註記：Prior period 前一期 = 起始日之前多抓的一期，供計算第一期報酬",
         "Flag: Prior period = one extra period before the start date, for computing the first return"),
        ("Flag 註記：Partial 未完整 = 結束日落在週期中間，該期資料不完整（仍保留最後交易日的數字）",
         "Flag: Partial = the end date falls inside the period; the period is incomplete (the last trading day's value is kept)"),
        ("Flag 註記：Intraday 盤中 = 日資料最後一筆為當日尚未收盤的盤中價",
         "Flag: Intraday = the last daily row is an unfinished intraday price"),
    ]
    start = info.max_row + 2
    for i, (zh, en) in enumerate(notes):
        info.cell(start + i * 2, 1, zh)
        info.cell(start + i * 2 + 1, 1, en)
    info.cell(start, 1).font = Font(bold=True, color="C00000")
    info.cell(start, 1).fill = SAMPLE_FILL

    wb.save(OUT)
    print("saved:", OUT)


if __name__ == "__main__":
    main()
